# src/collector.py
from __future__ import annotations
import os, sys, time, errno, msvcrt, threading, requests, openpyxl
from requests.exceptions import ChunkedEncodingError, ConnectionError, Timeout
try:
    from urllib3.exceptions import ProtocolError, IncompleteRead
except ImportError:  # pragma: no cover
    class ProtocolError(Exception):
        pass

    class IncompleteRead(Exception):
        pass
from dotenv import load_dotenv
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

load_dotenv()
API_TOKEN = os.getenv("SKETCHFAB_TOKEN")
HEADERS = {"Authorization": f"Token {API_TOKEN}"} if API_TOKEN else {}

ALIASES = {"girl": "Female", "legend of zelda": "Zelda", "links awakening": "Zelda"}
SINGLE_ASSIGNMENT_COLLECTIONS = {"hands", "gauntlets", "feet", "shoes"}

DATA_DIR = "data"
os.makedirs(DATA_DIR, exist_ok=True)
XL_PATH = os.path.join(DATA_DIR, "sketchfab_data.xlsx")

def _check_file_not_open(filepath: str) -> None:
    try:
        with open(filepath, "r+b") as f:
            msvcrt.locking(f.fileno(), msvcrt.LK_NBLCK, 1)
            msvcrt.locking(f.fileno(), msvcrt.LK_UNLCK, 1)
    except OSError as e:
        if e.errno == errno.EACCES:
            print(f"X '{filepath}' is open. Close it and rerun."); sys.exit(1)

def _auto_format(ws) -> None:
    """Cheap column widths — sample header + first rows only (full scan is huge on Collect)."""
    max_row = min(ws.max_row or 1, 40)
    max_col = ws.max_column or 1
    for c in range(1, max_col + 1):
        max_len = 8
        for r in range(1, max_row + 1):
            val = ws.cell(row=r, column=c).value
            if val is None:
                continue
            max_len = max(max_len, min(len(str(val)), 58))
        ws.column_dimensions[get_column_letter(c)].width = min(max_len + 2, 60)
    ws.freeze_panes = "A2"

def _remove_trailing_empty(ws) -> None:
    while ws.max_row > 1 and all(cell.value in (None, "") for cell in ws[ws.max_row]):
        ws.delete_rows(ws.max_row)
    while ws.max_column > 1 and all(ws.cell(row=r, column=ws.max_column).value in (None, "") for r in range(1, ws.max_row + 1)):
        ws.delete_cols(ws.max_column)

def _thumbnail_url(model: dict, max_width: int = 200) -> str:
    """Prefer a small thumbnail for fast grid loading (not the largest available)."""
    thumbs = model.get("thumbnails") or {}
    images = thumbs.get("images") or []
    if not images:
        return ""
    suitable = [i for i in images if (i.get("width") or 0) >= 64]
    pool = suitable or images
    in_range = [i for i in pool if (i.get("width") or 9999) <= max_width]
    pick_from = in_range or pool
    best = min(pick_from, key=lambda i: (i.get("width") or 0) * (i.get("height") or 0))
    return best.get("url") or ""


def _thumbnail_url_hd(model: dict, min_width: int = 512) -> str:
    """Larger thumbnail for LOD upgrade after low-res paints (sharp at dense grid sizes)."""
    thumbs = model.get("thumbnails") or {}
    images = thumbs.get("images") or []
    if not images:
        return ""
    suitable = [i for i in images if (i.get("width") or 0) >= min_width]
    pool = suitable or images
    best = max(pool, key=lambda i: (i.get("width") or 0) * (i.get("height") or 0))
    return best.get("url") or ""


def _merge_name_lists(*cells: object) -> list[str]:
    out: list[str] = []
    seen: set[str] = set()
    for cell in cells:
        raw = "" if cell is None else str(cell)
        for part in raw.split(","):
            name = part.strip()
            if not name:
                continue
            key = name.casefold()
            if key in seen:
                continue
            seen.add(key)
            out.append(name)
    return out


def _merge_already_in(
    prev: str,
    live_names: list[str],
    *,
    crawled: set[str],
    known: set[str],
) -> str:
    """
    Rebuild Already In from live membership without wiping unfinished crawls.

    - Collections we successfully re-listed: use live membership only.
    - Collections we failed / skipped (throttle, cancel): keep previous names.
    - Collections that no longer exist on the account: drop.
    - Local "Unlisted" shelf (private models Sketchfab won't collect): always keep,
      plus any other prev labels on that row (so N/W filter collections survive).
    """
    crawled_cf = {c.casefold() for c in crawled if c}
    known_cf = {k.casefold() for k in known if k}
    live = [n.strip() for n in live_names if str(n).strip()]
    live_cf = {n.casefold() for n in live}
    prev_parts = _merge_name_lists(prev)
    has_unlisted = any(p.casefold() == "unlisted" for p in prev_parts)
    kept: list[str] = []
    for name in prev_parts:
        key = name.casefold()
        if key == "unlisted":
            kept.append(name)
            continue
        if key in crawled_cf:
            # Live list is authoritative — unless this is an Unlisted shelf row
            # keeping intended collection labels for N/W filters.
            if has_unlisted and key not in live_cf:
                kept.append(name)
            continue
        if key not in known_cf:
            if has_unlisted:
                kept.append(name)
            continue
        kept.append(name)
    return ", ".join(_merge_name_lists(", ".join(kept), ", ".join(live)))


def _requeue_unconfirmed(
    prev_already: str,
    prev_assigned: str,
    live_names: list[str],
    *,
    crawled: set[str],
) -> str:
    """
    If we previously believed a model was in a collection (Already In after Push)
    but a successful crawl does not list it there, put that name back into Assigned
    so it shows as Pending again instead of vanishing as Unassigned.
    """
    live_cf = {n.strip().casefold() for n in live_names if str(n).strip()}
    crawled_cf = {c.casefold() for c in crawled if c}
    prev_already_parts = _merge_name_lists(prev_already)
    has_unlisted = any(p.casefold() == "unlisted" for p in prev_already_parts)
    requeue: list[str] = []
    for name in prev_already_parts:
        key = name.casefold()
        if key == "unlisted":
            continue  # local shelf — never push as Assigned
        if has_unlisted:
            continue  # keep filter labels on Unlisted shelf; don't re-pend
        if key in crawled_cf and key not in live_cf:
            requeue.append(name)
    # Keep existing Assigned, drop names confirmed live (no longer need to push).
    keep_assigned: list[str] = []
    for name in _merge_name_lists(prev_assigned):
        if name.casefold() in live_cf:
            continue
        keep_assigned.append(name)
    return ", ".join(_merge_name_lists(", ".join(keep_assigned), ", ".join(requeue)))


def _load_workbook_liked_map(path: str) -> dict[str, dict]:
    """Read preserve-fields map from an xlsx path (used for main + .bak repair)."""
    if not path or not os.path.exists(path):
        return {}
    # openpyxl rejects .bak — copy to a temp xlsx when needed.
    src = path
    tmp = None
    if path.lower().endswith(".bak"):
        tmp = path + ".__repair__.xlsx"
        try:
            import shutil
            shutil.copy2(path, tmp)
            src = tmp
        except OSError:
            return {}
    try:
        wb = openpyxl.load_workbook(src, read_only=True, data_only=True)
        if "Liked Models" not in wb.sheetnames:
            wb.close()
            return {}
        ws = wb["Liked Models"]
        rows_iter = ws.iter_rows(min_row=1, values_only=True)
        headers = list(next(rows_iter))
        def idx(name):
            return headers.index(name) if name in headers else -1

        fields = {
            "assigned": idx("Assigned Collection") if idx("Assigned Collection") >= 0 else idx("Assigned Collection(s)"),
            "already_in": idx("Already In Collection(s)"),
            "auto": idx("Auto-Assigned Collection(s)"),
            "suggested": idx("Suggested Collection(s)"),
            "fuzzy": idx("Fuzzy Matched Collection(s)") if idx("Fuzzy Matched Collection(s)") >= 0 else idx("Fuzzy Match Collection(s)"),
            "manual": idx("Manual"),
            "push_sent": idx("Push Sent"),
            "pushed_at": idx("Pushed At"),
            "assigned_at": idx("Assigned At"),
            "collected_at": idx("Collected At"),
            "liked_order": idx("Liked Order"),
            "liked_at": idx("Liked At"),
            "author_username": idx("Author Username"),
            "downloaded": idx("Downloaded"),
            "download_path": idx("Download Path"),
            "thumbnail": idx("Thumbnail"),
            "thumbnail_hd": idx("Thumbnail HD"),
            "notes": idx("Assignment Notes"),
            "original_source": idx("Original Source"),
            "original_format": idx("Original Format"),
            "original_size": idx("Original Size"),
        }
        uid_i = idx("UID") if idx("UID") >= 0 else idx("Model UID")
        out: dict[str, dict] = {}
        for row in rows_iter:
            uid = row[uid_i] if uid_i >= 0 else None
            if not uid:
                continue
            out[str(uid)] = {
                k: (row[i] if i >= 0 and i < len(row) else "") for k, i in fields.items()
            }
        wb.close()
        return out
    except Exception:
        return {}
    finally:
        if tmp and os.path.exists(tmp):
            try:
                os.remove(tmp)
            except OSError:
                pass


def _load_preserved_rows(path=XL_PATH) -> dict[str, dict]:
    """Load per-UID values we should not wipe on re-collect.

    Also repairs rows that Push marked as sent then Collect wiped: if Push Sent is set
    but Assigned/Already In are empty, restore collection names from the .bak into
    Assigned so they reappear in Pending Push.
    """
    out = _load_workbook_liked_map(path)
    bak = (path or "") + ".bak"
    if not out or not os.path.exists(bak):
        return out
    bak_rows = _load_workbook_liked_map(bak)
    if not bak_rows:
        return out
    repaired = 0
    for uid, row in out.items():
        push = str(row.get("push_sent") or "").strip().lower() in {"yes", "y", "true", "1"}
        assigned = str(row.get("assigned") or "").strip()
        already = str(row.get("already_in") or "").strip()
        if not push or assigned or already:
            continue
        br = bak_rows.get(uid) or {}
        restore = str(br.get("already_in") or br.get("assigned") or br.get("manual") or "").strip()
        if not restore:
            continue
        # These never landed (or were wiped) — queue them for Push again.
        row["assigned"] = restore
        repaired += 1
    if repaired:
        _stamp(f"Restored Assigned on {repaired:,} Push-Sent row(s) from workbook backup (re-queue).")
    return out

def _load_assigned_collections(path=XL_PATH) -> dict[str, list[str]]:
    assigned = {}
    if not os.path.exists(path): return assigned
    wb = openpyxl.load_workbook(path)
    if "Liked Models" not in wb.sheetnames: return assigned
    for row in wb["Liked Models"].iter_rows(min_row=2, values_only=True):
        uid, assigned_str = row[1], row[2]
        if uid and assigned_str:
            assigned[uid] = [c.strip() for c in str(assigned_str).split(",") if c.strip()]
    return assigned

_HTTP_LOCAL = threading.local()
_GET_RETRIES = 10
_GET_TIMEOUT = (15, 120)  # connect, read — likes pages can be large
_RETRY_EXC = (ChunkedEncodingError, ConnectionError, Timeout, ProtocolError, IncompleteRead, OSError)
_RETRY_HTTP = frozenset({429, 500, 502, 503, 504})


def _session() -> requests.Session:
    """Per-thread session — shared Session is not safe under collection crawl workers."""
    sess = getattr(_HTTP_LOCAL, "session", None)
    if sess is None:
        sess = requests.Session()
        sess.headers.update({
            "Authorization": f"Token {API_TOKEN}",
            "Accept": "application/json",
            "User-Agent": "Sketchfab-Collections-Pipeline/5.0",
        })
        _HTTP_LOCAL.session = sess
    return sess


def _get(url: str, *, quiet: bool = False, max_retries: int | None = None) -> requests.Response:
    if not API_TOKEN:
        raise RuntimeError("SKETCHFAB_TOKEN missing (set it in .env).")
    sess = _session()
    retries = max_retries if max_retries is not None else _GET_RETRIES
    last_err: Exception | None = None
    for attempt in range(retries):
        try:
            r = sess.get(url, timeout=_GET_TIMEOUT)
            if r.status_code in _RETRY_HTTP and attempt + 1 < retries:
                wait = min(3.0 * (2 ** attempt), 45.0)
                retry_after = r.headers.get("Retry-After")
                if retry_after:
                    try:
                        wait = max(wait, float(retry_after))
                    except ValueError:
                        pass
                if not quiet:
                    print(
                        f"  API HTTP {r.status_code} — retry {attempt + 1}/{retries - 1} in {wait:.0f}s"
                    )
                time.sleep(wait)
                continue
            r.raise_for_status()
            return r
        except _RETRY_EXC as e:
            last_err = e
            if attempt + 1 >= retries:
                break
            wait = min(3.0 * (2 ** attempt), 45.0)
            if not quiet:
                print(
                    f"  Network glitch on API page — retry {attempt + 1}/{retries - 1} "
                    f"in {wait:.0f}s ({type(e).__name__})"
                )
            time.sleep(wait)
            try:
                sess.close()
            except Exception:
                pass
            _HTTP_LOCAL.session = None
            sess = _session()
    if last_err:
        raise last_err
    raise RuntimeError("API request failed")

# Cache collection→models during a Collect so the Collections sheet doesn't re-crawl.
_COLLECTION_MODELS_CACHE: dict[str, list[dict]] = {}
_STAMP_REPLACE_ACTIVE = False
_STAMP_T0: float | None = None


def _stamp(msg: str, *, replace: bool = False) -> None:
    """Print a status line. replace=True overwrites the previous replace line (CMD progress)."""
    global _STAMP_REPLACE_ACTIVE, _STAMP_T0
    now = time.strftime("%H:%M:%S")
    if _STAMP_T0 is None:
        _STAMP_T0 = time.monotonic()
    elapsed = time.monotonic() - _STAMP_T0
    if elapsed < 60:
        timing = f"{elapsed:.0f}s"
    else:
        timing = f"{int(elapsed // 60)}m {int(elapsed % 60):02d}s"
    text = f"{now} [{timing}] {msg}"
    if replace:
        # Pad to clear leftovers from longer previous lines.
        print(f"\r{text:<120}", end="", flush=True)
        _STAMP_REPLACE_ACTIVE = True
    else:
        if _STAMP_REPLACE_ACTIVE:
            print(flush=True)
            _STAMP_REPLACE_ACTIVE = False
        print(text, flush=True)


def _stamp_reset_timer() -> None:
    global _STAMP_T0, _STAMP_REPLACE_ACTIVE
    _STAMP_T0 = time.monotonic()
    if _STAMP_REPLACE_ACTIVE:
        print(flush=True)
        _STAMP_REPLACE_ACTIVE = False


def get_likes(
    should_cancel=None,
    on_progress=None,
    *,
    max_pages: int | None = None,
    target_count: int | None = None,
    wb_uid_count: int = 0,
    quiet: bool = False,
) -> list[dict]:
    likes, url, page = [], "https://api.sketchfab.com/v3/me/likes?per_page=100&sort_by=-likedAt", 0
    probe_retries = 3 if quiet else None
    while url:
        if should_cancel and should_cancel():
            if not quiet:
                _stamp(f"Collect cancelled during likes fetch ({len(likes):,} so far).")
            break
        if max_pages is not None and page >= max_pages:
            if target_count is None:
                break
            if target_count is not None and len(likes) >= target_count:
                break
        r = _get(url, quiet=quiet, max_retries=probe_retries)
        data = r.json()
        batch = data.get("results", [])
        likes.extend(batch)
        page += 1
        if not quiet:
            msg = f"Likes… page {page} · {len(likes):,} total"
            _stamp(msg, replace=True)
            if on_progress:
                try:
                    on_progress(page, 0, msg)
                except Exception:
                    pass
        if target_count is not None and len(likes) >= target_count:
            break
        url = data.get("next")
        if url:
            time.sleep(0.05)
    if likes and not quiet:
        _stamp(f"Likes done — {len(likes):,} models.")
    return likes


def get_collections() -> list[dict]:
    cols, url, page = [], "https://api.sketchfab.com/v3/me/collections?per_page=100", 0
    while url:
        r = _get(url)
        data = r.json()
        batch = data.get("results", [])
        cols.extend(batch)
        page += 1
        _stamp(f"Listing collections… page {page} · {len(cols)}", replace=True)
        url = data.get("next")
        if url:
            time.sleep(0.05)
    return cols


def get_subscriptions() -> list[dict]:
    subs, url, page = [], "https://api.sketchfab.com/v3/me/subscriptions?per_page=100", 0
    while url:
        r = _get(url)
        data = r.json()
        batch = data.get("results", [])
        subs.extend(batch)
        page += 1
        _stamp(f"Subscriptions… page {page} · {len(subs)}", replace=True)
        url = data.get("next")
        if url:
            time.sleep(0.05)
    return subs


def get_models_in_collection(uid: str, *, use_cache: bool = True) -> list[dict]:
    # restricted=1 is required or age-restricted models are omitted from results
    # even when they belong to the collection (makes Push/Collect look like a wipe).
    uid = str(uid or "").strip()
    if use_cache and uid in _COLLECTION_MODELS_CACHE:
        return _COLLECTION_MODELS_CACHE[uid]
    models, url = [], f"https://api.sketchfab.com/v3/collections/{uid}/models?per_page=100&restricted=1"
    while url:
        r = _get(url)
        data = r.json()
        models.extend(data.get("results", []))
        url = data.get("next")
        if url:
            time.sleep(0.04)
    if use_cache and uid:
        _COLLECTION_MODELS_CACHE[uid] = models
    return models


def _model_uid_from_collection_item(m: dict) -> str:
    if not isinstance(m, dict):
        return ""
    uid = m.get("uid") or ""
    if not uid:
        nested = m.get("model")
        if isinstance(nested, dict):
            uid = nested.get("uid") or ""
        elif nested:
            uid = str(nested)
    return str(uid).strip()

def build_uid_to_collections_map(
    on_progress=None,
    should_cancel=None,
    cols: list[dict] | None = None,
    only_collection_uids: set[str] | None = None,
) -> tuple[dict[str, list[str]], set[str], set[str]]:
    """
    Returns (uid→collection names, successfully crawled names, all known collection names).
    If only_collection_uids is set, list models only for those collections (quick Collect).
    """
    from concurrent.futures import ThreadPoolExecutor, as_completed

    _stamp("Checking collections…")
    if cols is None:
        cols = get_collections()
    known = {(c.get("name") or "").strip() for c in cols if (c.get("name") or "").strip()}
    mapping: dict[str, list[str]] = {}
    crawled: set[str] = set()
    total = len(cols)
    done = 0
    cancelled = False

    def _one(col: dict) -> tuple[str, list[str], bool]:
        name = (col.get("name") or "").strip()
        uid = col.get("uid") or ""
        if should_cancel and should_cancel():
            return name, [], False
        if not uid:
            return name, [], False
        if only_collection_uids is not None and str(uid).strip() not in only_collection_uids:
            return name, [], False
        try:
            model_uids = [
                u for u in (_model_uid_from_collection_item(m) for m in get_models_in_collection(uid)) if u
            ]
            return name, model_uids, True
        except Exception as exc:
            _stamp(f"  collection failed: {name}: {exc}")
            return name, [], False

    workers = max(1, min(4, total or 1))
    with ThreadPoolExecutor(max_workers=workers) as pool:
        futs = [pool.submit(_one, c) for c in cols]
        for fut in as_completed(futs):
            if should_cancel and should_cancel():
                cancelled = True
                for f in futs:
                    f.cancel()
                break
            name, model_uids, ok = fut.result()
            if not ok:
                done += 1
                if on_progress:
                    try:
                        on_progress(done, total, f"Checking collections… {done}/{total} (skipped {name})")
                    except Exception:
                        pass
                _stamp(f"Checking collections… {done}/{total} (skipped {name})", replace=True)
                continue
            crawled.add(name)
            for mu in model_uids:
                mapping.setdefault(mu, []).append(name)
            done += 1
            msg = f"Checking collections… {done}/{total} · {name}"
            if on_progress:
                try:
                    on_progress(done, total, msg)
                except Exception:
                    pass
            _stamp(msg, replace=True)
    if cancelled:
        _stamp(f"Checking collections stopped — {len(crawled)}/{total} listed.")
    else:
        _stamp(f"Collections checked — {len(crawled)}/{total}.")
    return mapping, crawled, known

def _apply_aliases(text: str) -> str:
    for a, canon in ALIASES.items(): text = text.replace(a, canon)
    return text

def _fuzzy_match_collections(tags: list[dict], collections: list[str]) -> str:
    import difflib
    tag_names = [t["name"].lower() for t in tags]
    matched: set[str] = set()
    for tag in tag_names:
        for m in difflib.get_close_matches(tag, [c.lower() for c in collections], cutoff=0.7):
            for c in collections:
                if c.lower() == m: matched.add(c)
    return ", ".join(sorted(matched))

def _suggest_collections(name: str, tags: list[dict], collections: list[str]) -> str:
    name = _apply_aliases((name or "").lower())
    tag_names = [_apply_aliases(t["name"].lower()) for t in tags]
    hits = [c for c in collections if c.lower() in name or any(c.lower() in t for t in tag_names)]
    return ", ".join(hits)

def _auto_assign(tags_str: str, suggested_str: str, fuzzy_str: str) -> str:
    tags      = [t.strip().lower() for t in (tags_str or "").split(",") if t.strip()]
    suggested = [s.strip().lower() for s in (suggested_str or "").split(",") if s.strip()]
    fuzzy     = [f.strip().lower() for f in (fuzzy_str or "").split(",") if f.strip()]
    all_cands = set(tags + suggested + fuzzy)
    votes = {c: sum([c in tags, c in suggested, c in fuzzy]) for c in all_cands}
    strong = [c for c, n in votes.items() if n == 3]
    singles = [c for c in strong if c in SINGLE_ASSIGNMENT_COLLECTIONS]
    if len(singles) == 1: return singles[0]
    if len(singles) > 1:  return ""
    return ", ".join(sorted(strong))

def _get_collection_names(cols: list[dict] | None = None) -> list[str]:
    src = cols if cols is not None else get_collections()
    return [(c.get("name") or "").strip() for c in src if (c.get("name") or "").strip()]


def _write_liked_models_sheet(
    wb: Workbook,
    likes,
    uid2cols,
    assigned_map,
    col_names,
    *,
    crawled: set[str] | None = None,
    known: set[str] | None = None,
) -> None:
    ws = wb.active; ws.title = "Liked Models"
    headers = [
        "Name", "UID", "Thumbnail", "Thumbnail HD", "Liked Order", "Liked At",
        "Assigned Collection(s)", "Already In Collection(s)",
        "Auto-Assigned Collection(s)", "Suggested Collection(s)", "Fuzzy Match Collection(s)",
        "Tags", "Categories", "Author", "Author Username", "License", "Face Count", "Views", "Likes", "Downloads", "Downloadable",
        "Original Source", "Original Format", "Original Size",
        "Downloaded", "Download Path",
        "Manual", "Push Sent", "Pushed At", "Assigned At", "Collected At", "Assignment Notes",
    ]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        if "Collection" in (cell.value or ""):
            cell.fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")

    prev_rows = _load_preserved_rows()
    from datetime import datetime, timezone
    now_iso = datetime.now(timezone.utc).replace(microsecond=0).isoformat()
    collect_stamp = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")
    crawled_names = set(crawled or [])
    known_names = set(known or []) or set(col_names or [])
    for order, m in enumerate(likes, start=1):
        name = m.get("name"); uid = m.get("uid"); tags = m.get("tags", [])
        tag_str = ", ".join(t["name"] for t in tags)
        suggested = _suggest_collections(name, tags, col_names)
        fuzzy = _fuzzy_match_collections(tags, col_names)
        auto = _auto_assign(tag_str, suggested, fuzzy)

        prev = prev_rows.get(str(uid), {})
        liked_at_prev = (prev.get("liked_at") or "").strip()
        if liked_at_prev:
            liked_at_val = liked_at_prev
        elif str(uid) in prev_rows:
            liked_at_val = ""
        else:
            liked_at_val = now_iso
        assigned_val = _requeue_unconfirmed(
            prev.get("already_in", ""),
            prev.get("assigned", ""),
            uid2cols.get(uid, []) or [],
            crawled=crawled_names,
        )
        # Never blank out Push membership just because this Collect's crawl missed a collection.
        already_in = _merge_already_in(
            prev.get("already_in", ""),
            uid2cols.get(uid, []) or [],
            crawled=crawled_names,
            known=known_names,
        )
        suggested_val = prev.get("suggested") or suggested
        fuzzy_val = prev.get("fuzzy") or fuzzy
        auto_val = prev.get("auto") or auto
        thumb = _thumbnail_url(m) or prev.get("thumbnail", "")
        thumb_hd = _thumbnail_url_hd(m) or prev.get("thumbnail_hd", "") or thumb

        lic = m.get("license") or {}
        # Don't .strip("()") — that strips trailing ')' and truncates e.g. "CC Attribution (by)".
        _lic_label = str(lic.get("label") or "").strip()
        _lic_slug = str(lic.get("slug") or "").strip()
        if _lic_label and _lic_slug:
            lic_str = f"{_lic_label} ({_lic_slug})"
        else:
            lic_str = _lic_label or _lic_slug
        is_dl = m.get("isDownloadable", False)
        user = m.get("user") or {}
        author = user.get("displayName") or user.get("username") or ""
        author_username = str(user.get("username") or "").strip()
        faces = m.get("faceCount")
        face_str = f"{int(faces):,}" if faces is not None else ""
        views = m.get("viewCount")
        likes_n = m.get("likeCount")
        dls = m.get("downloadCount")
        views_str = f"{int(views):,}" if views is not None else ""
        likes_str = f"{int(likes_n):,}" if likes_n is not None else ""
        dls_str = f"{int(dls):,}" if dls is not None else ""
        cats = m.get("categories") or []
        if isinstance(cats, list):
            cat_str = ", ".join(
                (c.get("name") if isinstance(c, dict) else str(c)) for c in cats
            )
        else:
            cat_str = ""
        url = m.get("viewerUrl") or f"https://sketchfab.com/3d-models/{uid}"
        # Keep Assigned At only while something is still pending push.
        assigned_at_val = prev.get("assigned_at", "") if str(assigned_val or "").strip() else ""

        ws.append([
            name, uid, thumb, thumb_hd, order, liked_at_val,
            assigned_val, already_in, auto_val, suggested_val, fuzzy_val,
            tag_str, cat_str, author, author_username or prev.get("author_username", ""), lic_str, face_str, views_str, likes_str, dls_str, "Yes" if is_dl else "No",
            prev.get("original_source", ""), prev.get("original_format", ""), prev.get("original_size", ""),
            prev.get("downloaded", ""), prev.get("download_path", ""),
            prev.get("manual", ""), prev.get("push_sent", ""), prev.get("pushed_at", ""),
            assigned_at_val,
            collect_stamp,
            prev.get("notes", ""),
        ])
        ws.cell(row=ws.max_row, column=1).hyperlink = url

    _auto_format(ws)

def _write_collections_sheet(wb: Workbook, cols: list[dict] | None = None) -> None:
    """Write Collections sheet using the Collect-time model cache (no second crawl)."""
    ws = wb.create_sheet("Collections")
    ws.append(["Collection Name", "Collection UID", "Slug", "Model Count", "Thumbnail", "Model Names"])
    for col in (cols if cols is not None else get_collections()):
        uid = str(col.get("uid") or "").strip()
        models = get_models_in_collection(uid, use_cache=True) if uid else []
        names = sorted([(m.get("name") or "") for m in models if m.get("name")])
        slug = col.get("slug") or ""
        thumb = ""
        # Prefer API collection thumbnails, else first model's thumb.
        cthumbs = (col.get("thumbnails") or {}).get("images") or []
        if cthumbs:
            by_w = sorted(cthumbs, key=lambda i: int(i.get("width") or 0))
            mid = by_w[len(by_w) // 2] if by_w else {}
            thumb = str(mid.get("url") or by_w[-1].get("url") or "")
        if not thumb and models:
            thumb = _thumbnail_url(models[0]) or _thumbnail_url_hd(models[0]) or ""
        ws.append([col.get("name", ""), uid, slug, len(names), thumb, ", ".join(names)])
        from collection_urls import collection_public_url
        ws.cell(row=ws.max_row, column=1).hyperlink = collection_public_url(
            col.get("name", ""), uid, slug=slug
        )
    _auto_format(ws)

def _write_subscriptions_sheet(wb: Workbook, subs: list[dict] | None = None) -> None:
    ws = wb.create_sheet("Subscribed Collections")
    ws.append(["Collection Name", "Collection UID", "Owner", "Owner Profile", "Model Count"])
    for sub in subs if subs is not None else get_subscriptions():
        user = sub.get("user") or {}
        ws.append([
            sub.get("name", ""),
            sub.get("uid", ""),
            user.get("displayName") or user.get("username") or "",
            user.get("profileUrl") or "",
            sub.get("modelCount", ""),
        ])
        uid = sub.get("uid")
        if uid:
            ws.cell(row=ws.max_row, column=1).hyperlink = f"https://sketchfab.com/collections/{uid}"
    _auto_format(ws)

def refresh_subscriptions_sheet(path: str = XL_PATH) -> int:
    """Fetch /me/subscriptions and update only the Subscribed Collections sheet."""
    if not os.path.exists(path):
        raise FileNotFoundError(path)
    subs = get_subscriptions()
    wb = openpyxl.load_workbook(path)
    if "Subscribed Collections" in wb.sheetnames:
        wb.remove(wb["Subscribed Collections"])
    _write_subscriptions_sheet(wb, subs)
    wb.save(path)
    print(f"Updated subscriptions sheet ({len(subs)} rows).")
    return len(subs)

def _existing_liked_count(path: str = XL_PATH) -> int:
    if not os.path.exists(path):
        return 0
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            if "Liked Models" not in wb.sheetnames:
                return 0
            # header + data rows
            return max(0, (wb["Liked Models"].max_row or 1) - 1)
        finally:
            wb.close()
    except Exception:
        return 0


def build_workbook(on_progress=None, should_cancel=None) -> str:
    _check_file_not_open(XL_PATH)
    _COLLECTION_MODELS_CACHE.clear()
    _stamp_reset_timer()
    prior_liked = _existing_liked_count()
    _stamp("Collect started.")
    if on_progress:
        on_progress(0, 0, "Fetching likes…")
    likes = get_likes(should_cancel=should_cancel, on_progress=on_progress)
    cancelled = bool(should_cancel and should_cancel())
    if cancelled and not likes:
        raise RuntimeError("Collect cancelled before likes finished.")
    # Never replace a large Liked sheet with a truncated fetch (Stop mid-Collect used to
    # rewrite 6k+ likes down to a few dozen and eat local Assigned rows with them).
    if prior_liked >= 100 and len(likes) < max(50, int(prior_liked * 0.85)):
        raise RuntimeError(
            f"Refusing to overwrite Liked sheet: workbook has {prior_liked:,} likes but "
            f"this Collect only fetched {len(likes):,}"
            + (" (stopped early — click Collect again and let likes finish)" if cancelled else "")
            + ". Your Sketchfab likes are still online; local file was left unchanged."
        )
    if on_progress:
        on_progress(0, 0, f"Likes done — checking collections…")
    # One collections list for map + sheet + names (avoid 2–3 extra list crawls).
    cols = get_collections()
    uid2cols, crawled, known = build_uid_to_collections_map(
        on_progress=on_progress,
        should_cancel=should_cancel,
        cols=cols,
    )
    if should_cancel and should_cancel():
        _stamp(
            f"Collect stopped after likes — writing workbook (keeping prior Already In for "
            f"{len(known) - len(crawled)} unlisted collection(s))…"
        )
    assigned = _load_assigned_collections()
    col_names = _get_collection_names(cols)

    wb = Workbook()
    if on_progress:
        on_progress(0, 0, "Writing workbook…")
    t_write = time.monotonic()
    _stamp("Writing workbook…")
    _write_liked_models_sheet(
        wb, likes, uid2cols, assigned, col_names, crawled=crawled, known=known
    )
    _write_collections_sheet(wb, cols)
    _write_subscriptions_sheet(wb)
    wb.save(XL_PATH)
    _COLLECTION_MODELS_CACHE.clear()
    _stamp(f"Saved {XL_PATH} (write {time.monotonic() - t_write:.1f}s)")
    return XL_PATH

def build_workbook_quick(
    on_progress=None,
    should_cancel=None,
    *,
    max_like_pages: int = 20,
    drift_uids: set[str] | None = None,
) -> str:
    """Recent Collect — smart likes fetch + targeted collection re-list when counts drift."""
    from sync_probe import fetch_account_counts, probe_collection_drifts
    from data_io import read_workbook

    _check_file_not_open(XL_PATH)
    _COLLECTION_MODELS_CACHE.clear()
    _stamp_reset_timer()
    prior_liked = _existing_liked_count()
    _stamp("Quick Collect started.")
    acct = fetch_account_counts()
    target_likes = acct.like_count

    if on_progress:
        on_progress(0, 0, "Fetching recent likes…")
    likes = get_likes(
        should_cancel=should_cancel,
        on_progress=on_progress,
        max_pages=max_like_pages,
        target_count=target_likes,
        wb_uid_count=prior_liked,
    )
    cancelled = bool(should_cancel and should_cancel())
    if cancelled and not likes:
        raise RuntimeError("Quick Collect cancelled before likes finished.")
    if target_likes is not None and len(likes) < target_likes and not cancelled:
        if on_progress:
            on_progress(0, 0, f"Fetching remaining likes ({len(likes):,}/{target_likes:,})…")
        likes = get_likes(
            should_cancel=should_cancel,
            on_progress=on_progress,
            target_count=target_likes,
            wb_uid_count=prior_liked,
        )
        cancelled = bool(should_cancel and should_cancel())
    if prior_liked >= 100 and len(likes) < max(50, int(prior_liked * 0.85)):
        raise RuntimeError(
            f"Refusing quick overwrite: workbook has {prior_liked:,} likes but fetch got {len(likes):,}."
        )
    if on_progress:
        on_progress(0, 0, "Listing collections…")
    cols = get_collections()
    crawl_uids = set(drift_uids or ())
    try:
        _likes_df, colls_df, _ = read_workbook(scan_downloads=False)
        for d in probe_collection_drifts(colls_df):
            if d.uid:
                crawl_uids.add(d.uid)
    except Exception:
        pass
    if on_progress:
        on_progress(0, 0, f"Re-checking {len(crawl_uids) or 'no'} drifted collection(s)…")
    uid2cols, crawled, known = build_uid_to_collections_map(
        on_progress=on_progress,
        should_cancel=should_cancel,
        cols=cols,
        only_collection_uids=crawl_uids,
    )
    assigned = _load_assigned_collections()
    col_names = _get_collection_names(cols)
    wb = Workbook()
    if on_progress:
        on_progress(0, 0, "Writing workbook…")
    _write_liked_models_sheet(
        wb, likes, uid2cols, assigned, col_names, crawled=crawled, known=known
    )
    _write_collections_sheet(wb, cols)
    _write_subscriptions_sheet(wb)
    wb.save(XL_PATH)
    _COLLECTION_MODELS_CACHE.clear()
    _stamp(f"Quick Collect saved {XL_PATH} ({len(likes):,} likes, {len(crawl_uids):,} col re-crawl)")
    return XL_PATH

if __name__ == "__main__":
    import pathlib
    import shutil
    if "--fix-encoding" in sys.argv:
        root = pathlib.Path(__file__).resolve().parent.parent
        for pc in root.rglob("__pycache__"):
            shutil.rmtree(pc, ignore_errors=True)
        fixed = 0
        for p in list(root.rglob("*.py")) + list(root.rglob("*.bat")):
            try:
                b = p.read_bytes()
                if len(b) >= 2 and b[:2] == b"\xff\xfe":
                    text = b.decode("utf-16")
                elif b"\x00" in b:
                    text = b.decode("utf-16-le")
                else:
                    continue
                p.write_text(text, encoding="utf-8", newline="\n")
                fixed += 1
            except OSError:
                pass
        if fixed:
            print(f"Fixed {fixed} corrupted file(s).")
    elif "--gui" in sys.argv:
        main_py = os.path.join(os.path.dirname(__file__), "main.py")
        os.execv(sys.executable, [sys.executable, main_py])
    else:
        build_workbook()
